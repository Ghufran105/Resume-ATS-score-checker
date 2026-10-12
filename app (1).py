"""ATS Resume Checker - Streamlit app powered by Gemini 3.8 Flash."""

import hashlib
import io
import json
import os
import re

import streamlit as st

MODEL_ID = "gemini-3.8-flash"
MAX_CHARS = 30_000  # keeps the prompt small and fast

PROMPT = """You are an expert ATS (Applicant Tracking System) analyst and senior recruiter.
Analyze the resume below{jd_clause} and reply with ONLY one valid JSON object.
No markdown, no code fences, no text before or after the JSON.

Use exactly this schema:
{{
  "ats_score": <integer 0-100>,
  "summary": "<2-3 sentence honest overall verdict>",
  "category_scores": {{
    "formatting": <0-100>,
    "keywords": <0-100>,
    "experience_impact": <0-100>,
    "skills": <0-100>,
    "education": <0-100>,
    "readability": <0-100>
  }},
  "strengths": ["<short point>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "issue": "<what is wrong>", "fix": "<exactly how to fix it>"}}
  ],
  "missing_keywords": ["<keyword>", "..."],
  "rewrite_examples": [
    {{"before": "<weak line copied from the resume>", "after": "<stronger rewrite>"}}
  ]
}}

Rules:
- Be strict and honest. Do not inflate the score. Most resumes score 40-80.
- Give 4-8 improvements, ordered High to Low priority.
- Give up to 10 missing keywords and up to 3 rewrite examples.
- Only quote lines that really exist in the resume. Never invent experience.
{jd_rule}
RESUME:
\"\"\"
{resume}
\"\"\"
{jd_block}"""


# ----------------------------------------------------------------- helpers
def extract_text(file_name: str, data: bytes) -> str:
    """Return plain text from a PDF, DOCX or TXT upload."""
    name = file_name.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    if name.endswith(".docx"):
        from docx import Document

        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))
        return "\n".join(parts)
    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Upload a PDF, DOCX or TXT file.")


def build_prompt(resume_text: str, job_description: str = "") -> str:
    jd = job_description.strip()
    if jd:
        return PROMPT.format(
            jd_clause=" against the job description provided",
            jd_rule="- Judge keyword match against the job description; missing_keywords must come from it.\n",
            resume=resume_text[:MAX_CHARS],
            jd_block=f'\nJOB DESCRIPTION:\n"""\n{jd[:8000]}\n"""\n',
        )
    return PROMPT.format(
        jd_clause="",
        jd_rule="- No job description was given, so judge general ATS-readiness and common industry keywords.\n",
        resume=resume_text[:MAX_CHARS],
        jd_block="",
    )


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def parse_result(raw: str) -> dict:
    """Turn the model reply into a clean, validated dict. Raises ValueError."""
    text = (raw or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE)
    start, end = text.find("{"), text.rfind("}")
    if start == -1 or end <= start:
        raise ValueError("Model did not return JSON.")
    data = json.loads(text[start : end + 1])
    if not isinstance(data, dict) or "ats_score" not in data:
        raise ValueError("JSON is missing 'ats_score'.")

    cats = data.get("category_scores") or {}
    improvements = []
    for item in data.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append(
                {
                    "priority": str(item.get("priority", "Medium")).capitalize(),
                    "issue": str(item.get("issue", "")),
                    "fix": str(item.get("fix", "")),
                }
            )
    rewrites = [
        {"before": str(r.get("before", "")), "after": str(r.get("after", ""))}
        for r in (data.get("rewrite_examples") or [])
        if isinstance(r, dict)
    ]
    return {
        "ats_score": _clamp(data.get("ats_score")),
        "summary": str(data.get("summary", "")),
        "category_scores": {str(k): _clamp(v) for k, v in cats.items()},
        "strengths": [str(s) for s in (data.get("strengths") or [])],
        "improvements": improvements,
        "missing_keywords": [str(k) for k in (data.get("missing_keywords") or [])],
        "rewrite_examples": rewrites,
    }


def analyze(client, resume_text: str, job_description: str = "") -> dict:
    """Call Gemini 3.8 Flash; retry once if the reply is not valid JSON."""
    prompt = build_prompt(resume_text, job_description)
    last_error = None
    for _ in range(2):
        interaction = client.interactions.create(
            model=MODEL_ID,
            input=prompt,
            generation_config={"thinking_level": "medium"},
        )
        try:
            return parse_result(interaction.output_text)
        except (ValueError, json.JSONDecodeError) as exc:
            last_error = exc
    raise RuntimeError(f"Could not read the AI response: {last_error}")


def get_api_key() -> str:
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file locally
        key = ""
    return key or os.environ.get("GEMINI_API_KEY", "")


@st.cache_data(show_spinner=False, ttl=3600)
def cached_analyze(api_key: str, resume_text: str, jd: str, _digest: str) -> dict:
    from google import genai

    return analyze(genai.Client(api_key=api_key), resume_text, jd)


def score_label(score: int) -> str:
    if score >= 80:
        return "Strong"
    if score >= 60:
        return "Decent - needs work"
    return "Weak - big improvements needed"


# --------------------------------------------------------------------- UI
def main():
    st.set_page_config(page_title="ATS Resume Checker", page_icon="📄", layout="centered")
    st.title("📄 ATS Resume Checker")
    st.caption(f"Upload your resume, get an ATS score and clear fixes. Powered by {MODEL_ID}.")

    api_key = get_api_key()
    if not api_key:
        api_key = st.sidebar.text_input("Gemini API key", type="password")
        st.sidebar.markdown("Get a free key at [Google AI Studio](https://aistudio.google.com/apikey).")
    if not api_key:
        st.info("Add your Gemini API key (sidebar, or `GEMINI_API_KEY` in Streamlit secrets) to start.")

    uploaded = st.file_uploader("Upload resume", type=["pdf", "docx", "txt"])
    jd = st.text_area(
        "Job description (optional, improves keyword matching)",
        height=150,
        placeholder="Paste the job description here...",
    )

    if st.button("Analyze resume", type="primary", disabled=not (uploaded and api_key)):
        data = uploaded.getvalue()
        try:
            resume_text = extract_text(uploaded.name, data).strip()
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            st.stop()
        if len(resume_text) < 100:
            st.error(
                "Almost no text was found. If this is a scanned/image PDF, an ATS can't read it either. "
                "Export a text-based PDF or DOCX and try again."
            )
            st.stop()

        digest = hashlib.sha256(data + jd.encode()).hexdigest()
        try:
            with st.spinner("Analyzing your resume..."):
                result = cached_analyze(api_key, resume_text, jd.strip(), digest)
        except Exception as exc:
            st.error(f"Analysis failed: {exc}")
            st.stop()

        score = result["ats_score"]
        st.divider()
        col1, col2 = st.columns([1, 2])
        col1.metric("ATS score", f"{score}/100")
        col2.write(f"**{score_label(score)}**")
        col2.progress(score / 100)
        if result["summary"]:
            st.write(result["summary"])

        if result["category_scores"]:
            st.subheader("Score breakdown")
            for name, val in result["category_scores"].items():
                st.write(name.replace("_", " ").title())
                st.progress(val / 100, text=f"{val}/100")

        if result["strengths"]:
            st.subheader("✅ Strengths")
            for s in result["strengths"]:
                st.write(f"- {s}")

        if result["improvements"]:
            st.subheader("🛠️ Improvements")
            icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
            for item in result["improvements"]:
                with st.expander(f"{icons.get(item['priority'], '🟠')} {item['priority']}: {item['issue']}"):
                    st.write(item["fix"])

        if result["missing_keywords"]:
            st.subheader("🔑 Missing keywords")
            st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

        if result["rewrite_examples"]:
            st.subheader("✍️ Rewrite examples")
            for ex in result["rewrite_examples"]:
                st.markdown(f"**Before:** {ex['before']}")
                st.markdown(f"**After:** {ex['after']}")
                st.write("")

        st.download_button(
            "Download report (JSON)",
            json.dumps(result, indent=2),
            file_name="ats_report.json",
            mime="application/json",
        )

    st.caption("Your resume is sent to Google's Gemini API for analysis. Scores are AI estimates, not real ATS output.")


if __name__ == "__main__":
    main()
